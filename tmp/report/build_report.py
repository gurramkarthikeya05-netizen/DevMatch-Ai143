from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from zipfile import ZipFile
import ast, json, re
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(r'D:\DBMS_DevApp'); TMP=ROOT/'tmp/report'; OUT=ROOT/'output'; OUT.mkdir(exist_ok=True)
FRONT=Path(r'C:\Users\gurra\Downloads\DBMS Mini Project_ cover+certificate+Declaration+Abstract+acknowledgement+TOC.docx')
BODY=Path(r'C:\Users\gurra\Downloads\DBMS Mini Project _Report content format.docx')
APP=ROOT/'DevMatch_AI_DBMS_Mini_Project/devmatch-ai'; TITLE='DevMatch AI'
sources={str(f):sha256(f.read_bytes()).hexdigest() for f in (FRONT,BODY)}
front=Document(FRONT); bodyref=Document(BODY)
inventory={}
for f in (FRONT,BODY):
 with ZipFile(f) as z: inventory[str(f)]={n:{'size':len(z.read(n)),'sha256':sha256(z.read(n)).hexdigest()} for n in z.namelist()}
(TMP/'source-inventory.json').write_text(json.dumps(inventory,indent=2))
(TMP/'artifact.md').write_text('''# Report template contract
References and SHA256: '''+json.dumps(sources)+'''
Reference renders: tmp/report source named PDFs and page PNGs, exported using Word after the packaged renderer found no LibreOffice executable. Each reference has 8 pages and 1 section.
Front page system: Letter 21.59 x 27.94 cm; margins top/left/right 2.54 cm and bottom 2.1008 cm; black page border from original sectPr; lower Roman numbering starting i; cover suppresses footer. Keep logos, institutional text, degree, department, academic year and prescribed certificate/declaration wording. Personal names, USNs, guide, principal, HOD, dates and signatures become blank lines.
Body page system: A4 20.9991 x 29.6986 cm; top/bottom 1.8997 cm; left 3.1697 cm; right 2.54 cm; header/footer distance 1.27 cm. Preserve body header logo and source border and footer pattern. Restart Arabic numbering at 1.
Typography: Times New Roman; source body 12 pt justified 1.5 line spacing; chapter number/title 16 pt bold, title centered; subsections 14 pt bold. Front cover uses original 10/12/14/16/18 pt text and center alignment; certificate and declaration 12 pt justified 1.5. Code 9 pt Consolas and dictionary 10 pt are content components needed for code/table slots. Figure captions centered 10.5 pt below image, table captions above table.
Content order: cover, certificate, declaration, abstract <=100 words with 5 comma separated keywords, acknowledgement, contents, figures, tables; chapters 1 Introduction, 2 System Requirements, 3 Design, 4 Implementation, 5 Interpretation of Result, 6 Conclusion and Future Scope, References. Main content target 20 pages.
Slots: word/document.xml body paragraph indices 4,9-13,25,35,37-38,49,51,94,96,98 plus signature tables editable. TOC/list tables replace examples. Body reference chapter/subsection/caption elements are repeatable patterns; replace all instructional/example content and wearable figure. Original blank spacer paragraphs may be removed in expanded report and replaced with explicit page breaks to keep front matter on separate pages. Code excerpts may be formatted across lines without changing semantics. Source body header project title placeholder is editable.
Preserve-only: source logos, front border, geometry, institutional wording and original reference files. Source package inventories stored source-inventory.json. Merge requires two sections, imported body header/footer and body styles with explicit paragraph formatting. No comments/tracked changes retained. Source source parts are unchanged; final package necessarily merges relationships and expands pages.
Fidelity gates: logos unchanged, two page systems maintained, exact source hashes unchanged; all requested chapter/subsection slots filled; all personal slots blank; headings precede subsections with introductory prose; diagrams/screenshots referenced and captioned; static TOC page numbers verified against final Word render.
''',encoding='utf-8')

def font(size,bold=False):
 return ImageFont.truetype(str(Path(r'C:\Windows\Fonts')/('arialbd.ttf' if bold else 'arial.ttf')),size)
def diagram(name,kind):
 im=Image.new('RGB',(1600,1350),'white'); dr=ImageDraw.Draw(im)
 boxes={
 'users':(610,35,990,170),
 'developer_profiles':(100,270,500,445),'recruiter_profiles':(1100,270,1500,445),
 'skills':(30,600,330,760),'projects':(385,600,685,760),'jobs':(1040,600,1440,760),
 'project_technologies':(350,1000,720,1170),'job_skills':(1120,1000,1490,1170),
 'applications':(745,820,1080,990),'resume_analysis':(30,1000,330,1170),
 'application_status_history':(745,1160,1080,1320)}
 attrs={'users':['PK id','UK email, role'], 'developer_profiles':['PK id','FK UK user_id -> users.id'], 'recruiter_profiles':['PK id','FK UK user_id -> users.id'], 'skills':['PK id','FK developer_id, name'], 'projects':['PK id','FK developer_id, name'], 'jobs':['PK id','FK recruiter_id, title'], 'project_technologies':['PK id','FK project_id, name'], 'job_skills':['PK id','FK job_id, name'], 'applications':['PK id','FK developer_id, FK job_id'], 'resume_analysis':['PK id','FK developer_id, summary'], 'application_status_history':['PK id','FK application_id, status']}
 edges=[('users','developer_profiles','1 : 0..1'),('users','recruiter_profiles','1 : 0..1'),('developer_profiles','skills','1 : 0..N'),('developer_profiles','projects','1 : 0..N'),('developer_profiles','resume_analysis','1 : 0..N'),('developer_profiles','applications','1 : 0..N'),('recruiter_profiles','jobs','1 : 0..N'),('jobs','job_skills','1 : 0..N'),('jobs','applications','1 : 0..N'),('projects','project_technologies','1 : 0..N'),('applications','application_status_history','1 : 0..N')]
 routes={
 ('users','developer_profiles'):[(800,170),(800,220),(300,220),(300,270)],
 ('users','recruiter_profiles'):[(800,170),(800,220),(1300,220),(1300,270)],
 ('developer_profiles','skills'):[(220,445),(220,540),(180,540),(180,600)],
 ('developer_profiles','projects'):[(380,445),(380,540),(535,540),(535,600)],
 ('developer_profiles','resume_analysis'):[(100,390),(10,390),(10,930),(180,930),(180,1000)],
 ('developer_profiles','applications'):[(500,370),(770,370),(770,780),(860,780),(860,820)],
 ('recruiter_profiles','jobs'):[(1300,445),(1300,520),(1240,520),(1240,600)],
 ('jobs','job_skills'):[(1340,760),(1340,930),(1305,930),(1305,1000)],
 ('jobs','applications'):[(1040,680),(995,680),(995,820)],
 ('projects','project_technologies'):[(535,760),(535,1000)],
 ('applications','application_status_history'):[(900,990),(900,1160)]}
 for a,b,label in edges:
  ax,ay,ar,ab=boxes[a];bx,by,br,bb=boxes[b]
  start=((ax+ar)//2,ab);end=((bx+br)//2,by)
  if a=='developer_profiles' and b=='resume_analysis':start=(ax,ab-20);end=(bx,by+15)
  if a=='developer_profiles' and b=='applications':start=(ar,ab-30);end=(bx,by+35)
  if a=='jobs' and b=='applications':start=(ax,ab-20);end=(br,by+20)
  mid=(start[0],(start[1]+end[1])//2)
  route=routes[(a,b)];start=route[0];end=route[-1]
  dr.line(route,fill='#444444',width=4)
  dr.polygon([(end[0],end[1]),(end[0]-8,end[1]-15),(end[0]+8,end[1]-15)],fill='#444444')
  if kind=='er':
   x=route[-2][0]+12;y=(route[-2][1]+end[1])//2-15
   if (a,b)==('jobs','applications'):x=835;y=735
   dr.rectangle((x-5,y-3,x+155,y+30),fill='white');dr.text((x,y),label,font=font(24),fill='black')
 for k,(x,y,r,b) in boxes.items():
  dr.rectangle((x,y,r,b),fill='#f8f8f8',outline='black',width=3)
  title=k if kind=='schema' else {'developer_profiles':'Developer','recruiter_profiles':'Recruiter','application_status_history':'Status history','project_technologies':'Project technology','resume_analysis':'Resume analysis','job_skills':'Job requirement'}.get(k,k.title())
  fs=25 if len(title)>22 else 29
  dr.text((x+14,y+14),title,font=font(fs,True),fill='black')
  dr.line((x,y+60,r,y+60),fill='black',width=2)
  lines=attrs[k] if kind=='schema' else ['Identifier and attributes', 'Audit extension' if k=='application_status_history' else 'Relational entity']
  for i,t in enumerate(lines):dr.text((x+14,y+78+i*34),t,font=font(22),fill='black')
 im.save(TMP/(name+'.png'))
diagram('schema','schema');diagram('er','er')

# Start with the retained front matter and its original package assets.
doc=front
from docx.enum.style import WD_STYLE_TYPE
for name in ('Title','Heading 1','Heading 2','Heading 3'):
 if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
orig_paras=list(doc.paragraphs); orig_elements=list(doc._element.body)
def replace(p,text):
 base=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 for child in list(p._p):
  if child.tag!=qn('w:pPr'):p._p.remove(child)
 r=p.add_run(text)
 if base is not None:r._r.insert(0,base)
 for run in p.runs:run.font.highlight_color=None;run.font.color.rgb=RGBColor(0,0,0)
 return p
replacements={4:TITLE,9:'Name ____________________  USN ____________________',10:'Name ____________________  USN ____________________',11:'Name ____________________  USN ____________________',13:'Guide name ____________________  Designation ____________________',25:'This is to certify that the Database Management System Project work entitled "DevMatch AI" has been carried out by ____________________, USN ____________________, ____________________, USN ____________________, and ____________________, USN ____________________, bonafide students of CMR Institute of Technology, Bengaluru in partial fulfillment for the award of the Degree of Bachelor of Engineering in Electronics and Communication Engineering of the Visvesvaraya Technological University, Belagavi during the year 2026-27. It is certified that all corrections/suggestions indicated for the Internal Assessment have been incorporated in the report deposited in the departmental library. This Database Management System Project report has been approved as it satisfies the academic requirements in respect of project work prescribed for the said Degree.',35:'We, the students of VII semester from the Department of Electronics and Communication Engineering, CMR Institute of Technology, Bangalore declare that the project work entitled "DevMatch AI" has been successfully completed under the guidance of ____________________, designation ____________________, CMR Institute of Technology, Bengaluru. This project work is submitted in partial fulfillment of the requirements for the award of the Degree of Bachelor of Engineering in Electronics and Communication Engineering during the academic year 2026-2027. The matter embodied in the project report has not been submitted previously by anybody for the award of any degree or diploma to any university.',37:'Place: ____________________',38:'Date: ____________________',49:'DevMatch AI is a database driven developer portfolio and hiring platform. Developers manage profiles, skills, projects, resumes and applications, while recruiters publish jobs and review candidates. Flask and SQLAlchemy connect a responsive web interface to ten related MySQL tables. Normalized skill and technology records, integrity constraints, joins, views and a status audit trigger demonstrate DBMS concepts. PDF resume extraction supports transparent skill matching, with optional Gemini analysis. The system centralizes hiring information and allows recruiter review of automatically assigned application statuses.',51:'Keywords: MySQL, Flask, Database Normalization, Developer Portfolio, Skill Matching',94:'We thank ____________________, Principal, CMR Institute of Technology, Bengaluru, for encouragement and institutional support.',96:'We express our gratitude to ____________________, Head of the Department of Electronics and Communication Engineering, CMR Institute of Technology, Bengaluru, for support throughout the project.',98:'We thank our guide ____________________, designation ____________________, CMR Institute of Technology, Bengaluru, for guidance during the project work.'}
for i,t in replacements.items():replace(orig_paras[i],t)
orig_paras[4].style=doc.styles['Title'];orig_paras[4].alignment=WD_ALIGN_PARAGRAPH.CENTER
orig_paras[4].paragraph_format.space_after=Pt(6)
for r in orig_paras[4].runs:r.bold=True;r.font.size=Pt(16)
replace(orig_paras[92],'We express our gratitude to CMR Institute of Technology, Bengaluru, for providing the facilities and academic environment to carry out this Database Management System mini project.')
replace(orig_paras[100],'We thank our parents, friends and team members for their support and cooperation during the completion of this project.')
for t in doc.tables[:2]:
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    if 'Signature of Guide' in p.text:replace(p,'________________________\nSignature of Guide\nName ____________________\nDesignation ____________________')
    elif 'Signature of HOD' in p.text:replace(p,'________________________\nSignature of HOD\nName ____________________\nDesignation ____________________')
    elif 'NAME' in p.text:replace(p,'Name ____________________\nUSN ____________________')
    elif p.text=='Signature':replace(p,'____________________')
doc.tables[0].cell(0,0).text='________________________\nSignature of Guide\nName ____________________\nDesignation ____________________'
doc.tables[0].cell(0,2).text='________________________\nSignature of HOD\nName ____________________\nDesignation ____________________'
for c in doc.tables[0].rows[0].cells:
 for p in c.paragraphs:
  for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11)

# Remove instructions and spacer-only paragraphs, while preserving logos and page components.
starts=[0,20,33,47,88,115,126,147]
for p in list(doc.paragraphs):
 idx=orig_paras.index(p) if p in orig_paras else -1
 # python-docx creates new wrappers, so use underlying element identity.
 idx=next((i for i,q in enumerate(orig_paras) if q._p is p._p),-1)
 if idx in (128,149) or (not p.text.strip() and not p._p.xpath('.//w:drawing | .//w:pict')):
  p._p.getparent().remove(p._p)
 elif idx in starts[1:]:p.paragraph_format.page_break_before=True
for idx in (47,88,115,126,147):
 p=orig_paras[idx];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(35);p.paragraph_format.space_after=Pt(20)
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(16);r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
for p in doc.paragraphs:
 for r in p.runs:
  r.font.highlight_color=None;r.font.color.rgb=RGBColor(0,0,0)
 if p.text and not p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=False

def table_style(t,size=10.5,widths=None):
 t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  h=trpr.find(qn('w:trHeight'))
  if h is not None:trpr.remove(h)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[j])
   c.vertical_alignment=1
   pr=c._tc.get_or_add_tcPr()
   mar=OxmlElement('w:tcMar')
   for side in ('top','left','bottom','right'):
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'45');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar)
   borders=OxmlElement('w:tcBorders')
   for side in ('top','left','bottom','right'):
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');pr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.0
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=(i==0);r.font.color.rgb=RGBColor(0,0,0);r.font.highlight_color=None
def fill_existing(t,rows):
 while len(t.rows)>1:t._tbl.remove(t.rows[-1]._tr)
 for j,v in enumerate(rows[0]):t.cell(0,j).text=v
 for row in rows[1:]:
  c=t.add_row().cells
  for j,v in enumerate(row):c[j].text=str(v)
 table_style(t,11,[14,2.5])
tocrows=[('Contents','Page No.'),('Certificate','ii'),('Declaration','iii'),('Abstract','iv'),('Acknowledgement','v'),('Table of Contents','vi'),('List of Figures','vii'),('List of Tables','viii'),('1 Introduction','1'),('1.1 Objectives and 1.2 Scope of the project','2'),('2 System Requirements','3'),('2.1 Front End Tools and 2.2 Back End Tools','3'),('3 Design','4'),('3.1 Schema Diagram','4'),('3.2 ER Diagram','5'),('3.3 Database Tables and Normalization','6–9'),('4 Implementation','10'),('4.1 Installation Steps','10'),('4.2 Front End Code','11'),('4.3 Back End Code and Modules','12–15'),('5 Interpretation of Result','16'),('5.1 Result Snapshots','16–18'),('6 Conclusion and Future Scope','19'),('References','20')]
fill_existing(doc.tables[2],tocrows)
fill_existing(doc.tables[3],[('Figure','Page No.'),('Fig 3.1 Relational schema','4'),('Fig 3.2 Entity relationships and cardinalities','5'),('Fig 5.1 Public home page','16'),('Fig 5.2 Developer dashboard','16'),('Fig 5.3 Job listing and skill match','17'),('Fig 5.4 Recruiter applicant review','17'),('Fig 5.5 MySQL jobs demonstration','18')])
fill_existing(doc.tables[4],[('Table','Page No.'),('Table 2.1 Development requirements','3'),('Table 3.1 Account and profile tables','6'),('Table 3.2 Portfolio tables','7'),('Table 3.3 Job tables','8'),('Table 3.4 Application and analysis tables','9'),('Table 5.1 Database verification results','18')])
doc.tables[2].rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next=True

# Add the content template's own page system and recurring header/footer.
sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
refsec=bodyref.sections[0]
for a in ('page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance'):setattr(sec,a,getattr(refsec,a))
for x in sec._sectPr.xpath('./w:pgBorders | ./w:titlePg | ./w:pgNumType'):sec._sectPr.remove(x)
num=OxmlElement('w:pgNumType');num.set(qn('w:fmt'),'decimal');num.set(qn('w:start'),'1');sec._sectPr.append(num)
sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
def copy_hf(source,target):
 for c in list(target._element):target._element.remove(c)
 for c in source._element:target._element.append(deepcopy(c))
 for old,rel in source.part.rels.items():
  if 'image' in rel.reltype:
   rid,_=target.part.get_or_add_image(BytesIO(rel.target_part.blob))
   for el in target._element.iter():
    for key,v in list(el.attrib.items()):
     if v==old and key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):el.set(key,rid)
copy_hf(refsec.header,sec.header);copy_hf(refsec.footer,sec.footer)
for p in sec.header.paragraphs:
 for t in p._p.xpath('.//w:t'):
  if 'Project' in (t.text or ''):t.text='DevMatch AI'
  elif (t.text or '').strip()=='Title':t.text=''
 for x in p._p.xpath('.//w:highlight'):x.getparent().remove(x)
# Retain footer border and replace space-aligned text with tabs within the same margins.
fp=sec.footer.paragraphs[0];ppr=deepcopy(fp._p.pPr);fp.clear();fp._p.remove(fp._p.pPr);fp._p.insert(0,ppr)
fp.style=doc.styles['Normal']
fp.paragraph_format.space_after=Pt(0);fp.paragraph_format.line_spacing=1.0
for el in fp._p.pPr.xpath('./w:tabs'):fp._p.pPr.remove(el)
tabs=OxmlElement('w:tabs')
for pos in (4513,9026,4680,9360):
 clear=OxmlElement('w:tab');clear.set(qn('w:val'),'clear');clear.set(qn('w:pos'),str(pos));tabs.append(clear)
fp._p.pPr.append(tabs)
from docx.enum.text import WD_TAB_ALIGNMENT
fp.paragraph_format.tab_stops.add_tab_stop(Cm(15.25),WD_TAB_ALIGNMENT.RIGHT)
for tx in ('Dept. of ECE, CMRIT   2026-2027\tPage ',):
 r=fp.add_run(tx);r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=True
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
for st in ['Heading 1','Heading 2','Heading 3','Title']:
 doc.styles[st].font.name='Times New Roman';doc.styles[st].font.color.rgb=RGBColor(0,0,0)

page_index=0
pending_break=False
def para(t,size=12,bold=False,align=WD_ALIGN_PARAGRAPH.JUSTIFY):
 global pending_break
 p=doc.add_paragraph();p.alignment=align;p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(7)
 if pending_break:p.paragraph_format.page_break_before=True;pending_break=False
 r=p.add_run(t);r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
 return p
def newpage(ch=None,title=None):
 global page_index,pending_break
 if page_index:pending_break=True
 page_index+=1
 if ch:
  p=para('Chapter '+str(ch),16,True,WD_ALIGN_PARAGRAPH.LEFT);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6)
  p=para(title.upper(),16,True,WD_ALIGN_PARAGRAPH.CENTER);p.style=doc.styles['Heading 1'];p.paragraph_format.space_after=Pt(14)
def heading(t):
 p=para(t,14,True,WD_ALIGN_PARAGRAPH.LEFT);p.style=doc.styles['Heading 2'];p.paragraph_format.line_spacing=1.15;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
 return p
def caption(t,above=False):
 p=para(t,10.5,False,WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_with_next=above
def picture(name,width=6):
 if name in ('home','developer','match','applicants','dbms'):width=4.5
 if name in ('home','developer'):name=name+'_short'
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
 r=p.add_run();r.add_picture(str(TMP/(name+'.png')),width=Inches(width))
def table(title,headers,rows,widths=None):
 caption(title,True);t=doc.add_table(rows=1,cols=len(headers))
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 table_style(t,10,widths);return t
def code(t):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.0
 for i,line in enumerate(t.strip('\n').splitlines()):
  r=p.add_run(('\n' if i else '')+line);r.font.name='Consolas';r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
 return p

# 1 and 2 Introduction
newpage(1,'Introduction')
para('DevMatch AI is a web application that brings developer portfolios, job postings and application tracking into one relational database. A developer can present skills and projects, upload a resume, search jobs and monitor applications. A recruiter can maintain a company profile, publish vacancies and examine the evidence behind each candidate’s skill match.')
para('The project addresses the difficulty of keeping hiring information consistent when profiles, resumes, vacancies and application decisions are handled separately. Repeated candidate details and unstructured skill lists make searching and comparison difficult. DevMatch AI stores related information in dedicated MySQL tables and links records through primary and foreign keys.')
para('The system uses Flask for request handling, SQLAlchemy for database access and HTML, CSS and JavaScript for presentation. PDF text extraction provides resume evidence. Gemini can optionally produce a structured analysis, while a deterministic extraction path allows the resume workflow to function without an AI service. Supabase is used only for optional private resume file storage.')
para('The DBMS contribution is the design and implementation of ten application tables, normalized skill and project technology records, uniqueness constraints and relational queries. A separate SQL demonstration adds a view and an application status audit trigger. The report explains the requirements, design, implementation and observed local results of these components.')
para('Skill matching is intended to support review. The score measures coverage of a job’s recorded skill names, and it does not measure overall competence or predict hiring success. The recruiter can revise the initial application status after examining the portfolio and resume.')
newpage();heading('1.1 Objectives')
for t in ['Maintain developer and recruiter accounts with role specific access to their records.','Store portfolios, skills, jobs, applications and resume analyses in related MySQL tables.','Normalize multivalued job skills and project technologies to support searching and integrity.','Provide create, read, update and delete operations through the web interface.','Calculate a visible skill match using profile skills and the latest extracted resume skills.','Demonstrate keys, joins, filtering, grouping, views, transactions and a status history trigger.']:
 para('• '+t)
heading('1.2 Scope of the project')
para('The implemented scope includes registration and login, profile maintenance, developer portfolio management, job search, recruiter job management, one application per developer and job, PDF resume analysis, candidate discovery and application status review. The DBMS demonstration page displays current job records from the same MySQL database used by the application.')
para('The local implementation uses MySQL and can store resume files locally. Deployment configuration supports private Supabase Storage and shared Redis rate limiting. Interview scheduling, email notifications, payment processing and external job board integration are outside the current workflow. Scanned resumes require OCR support beyond the present text extraction path.')

# 3 requirements
newpage(2,'System Requirements')
para('The system requires a browser, a Python application environment and a reachable MySQL database. Optional services extend resume analysis and deployment; the core profile, job and application functions use the same relational backend in both local and hosted configurations.')
heading('2.1 Front End Tools')
para('HTML defines pages and forms; CSS provides responsive layouts, navigation, dashboards and job cards. Jinja templates bind server supplied records to the interface. JavaScript supports interactions such as copying job information and opening skill details. Form actions and database updates are processed by Flask rather than by browser access to MySQL.')
heading('2.2 Back End Tools')
para('Python and Flask implement routes, sessions and role checks. Flask SQLAlchemy defines models and database transactions; PyMySQL supplies the MySQL connection. PyMuPDF extracts PDF text, requests supports Gemini API calls, and python dotenv loads local configuration. Flask Limiter throttles selected account actions. MySQL Workbench or the command line client supports SQL demonstrations. References [1]–[6] identify the relevant official documentation.')
table('Table 2.1 Development requirements',['Component','Requirement or role'],[('Hardware','A computer capable of running Python, MySQL and a modern browser; 4 GB RAM and free storage are practical local starting points.'),('Software','Python 3.10 or later; MySQL 8; a browser; packages listed in requirements.txt.'),('Database','mysql+pymysql connection URL; dedicated database and account.'),('Optional services','Gemini API key; private Supabase resume bucket; Redis compatible limiter storage for production.')],[3.4,11.8])

# 4 schema
newpage(3,'Design')
para('The design separates shared account information from developer and recruiter data. Jobs belong to recruiters, while applications connect developers to jobs. Repeated skill and technology values are represented by child rows. Figure 3.1 summarizes the key links in the physical schema.')
heading('3.1 Schema Diagram')
picture('schema',5.7);caption('Fig 3.1 Relational schema')
para('PK denotes a primary key, FK a foreign key and UK a unique key. Ten tables are created by the application models. The eleventh table, application_status_history, is an audit extension installed by mysql_demo.sql. The unique developer and job pair in applications prevents duplicate applications.',11)
newpage();heading('3.2 ER Diagram')
para('Figure 3.2 shows the entity relationships and their cardinalities. Each profile references exactly one user, and a user has at most one profile of each kind. The application’s role logic selects the appropriate profile. Child entities can have zero or many records for their parent.',12)
picture('er',5.7);caption('Fig 3.2 Entity relationships and cardinalities')
para('The developer to job relationship is many to many and is resolved by Application. Each Application belongs to one Developer and one Job. Each Resume analysis and Skill belongs to one Developer. Each Job requirement belongs to one Job. The audit extension stores status transitions for an application.',11)

# 6-9 actual model dictionary parsed from source
tree=ast.parse((APP/'app.py').read_text(encoding='utf-8'))
models={}
for node in tree.body:
 if isinstance(node,ast.ClassDef):
  attrs=[];tn=None
  for n in node.body:
   if isinstance(n,ast.Assign):
    field=n.targets[0].id if isinstance(n.targets[0],ast.Name) else ''
    if field=='__tablename__':tn=ast.literal_eval(n.value)
    if isinstance(n.value,ast.Call) and ast.unparse(n.value.func)=='db.Column':
     v=n.value;dtype=ast.unparse(v.args[0]).replace('db.','');key=[]
     for a in v.args[1:]:
      if isinstance(a,ast.Call) and ast.unparse(a.func)=='db.ForeignKey':key.append('FK '+ast.literal_eval(a.args[0]))
     for kw in v.keywords:
      if kw.arg=='primary_key' and isinstance(kw.value,ast.Constant) and kw.value.value:key.append('PK')
      if kw.arg=='unique' and isinstance(kw.value,ast.Constant) and kw.value.value:key.append('UNIQUE')
     attrs.append((tn or node.name,field,dtype,', '.join(key) or 'Attribute'))
  if tn:models[tn]=attrs
def dictionary(names,title):
 rows=[]
 for name in names:rows+=models[name]
 table(title,['Table','Attribute','Type','Key or role'],rows,[4.1,3.8,2.8,4.5])
newpage();heading('3.3 Database Tables')
para('Tables 3.1–3.4 describe the actual SQLAlchemy model columns. VARCHAR lengths are shown as String lengths in the model. Integer identifiers are surrogate primary keys. The profile user_id columns are unique foreign keys, so a user cannot have duplicate records in the same profile table.',11)
dictionary(['users','developer_profiles','recruiter_profiles'],'Table 3.1 Account and profile tables')
para('Email is unique and the stored credential is a password hash. Phone and web links remain strings. User credentials are stored once and are referenced by the role specific profile.',11)
newpage();heading('3.3 Database Tables continued')
para('Portfolio records are separated into skills, projects and project technologies. The developer owns each skill and project, and each technology belongs to a project. This avoids storing a comma separated technology list in the projects table.',12)
dictionary(['skills','projects','project_technologies'],'Table 3.2 Portfolio tables')
para('The skills table has a uniqueness constraint on developer_id and name. Project technologies are unique on project_id and name. SQLAlchemy relationships provide access to these child rows, and form inputs are parsed into distinct entries before saving. The display properties can join values for presentation without reverting to a multivalued database column.',12)
newpage();heading('3.3 Database Tables continued')
para('Jobs contain vacancy details while job_skills stores one named requirement per row. The recruiter_id foreign key connects each vacancy to its owner. Table 3.3 lists the job columns and requirement attributes.',12)
dictionary(['jobs','job_skills'],'Table 3.3 Job tables')
para('The job_id and name pair is unique in job_skills. The current matching function compares every stored job skill name; importance and minimum_years are recorded but are not weighted in the score. The application model declares salary as String(100), while the SQL demonstration widens it to VARCHAR(150) using ALTER TABLE. This deliberate DDL change can therefore differ from the original model declaration.',12)
newpage();heading('3.3 Database Tables continued')
para('Applications resolve the developer and job association. Resume analyses hold extracted information used by the matching function. Status history is a separate SQL audit table, created only when the demonstration script is run.',11)
rows=models['applications']+models['resume_analysis']+[('application_status_history','id','INT','PK'),('application_status_history','application_id','INT','FK applications.id'),('application_status_history','old_status','VARCHAR(30)','Nullable'),('application_status_history','new_status','VARCHAR(30)','Required'),('application_status_history','changed_at','DATETIME','Timestamp')]
table('Table 3.4 Application and analysis tables',['Table','Attribute','Type','Key or role'],rows,[4.1,3.8,2.8,4.5])
heading('Normalization and integrity')
para('First normal form is supported by individual skill and technology rows. Second normal form keeps each non key fact associated with its complete entity or parent and value pair. Third normal form separates credentials, profiles, vacancies and application events. Resume text fields store summaries or extracted text as analysis payloads; the matching routine parses extracted skill text rather than treating it as a normalized skill relation. The unique application pair and foreign keys enforce important integrity rules.',11)

# 10 implementation
newpage(4,'Implementation')
para('The implementation connects browser forms to Flask routes, validates the request, reads or changes SQLAlchemy models and commits successful changes to MySQL. Jinja templates render the resulting records. The modules include accounts, portfolios, recruiter jobs, applications, resume analysis, skill matching and database demonstrations.')
heading('4.1 Installation Steps for Execution of the Project')
para('Create a virtual environment in the devmatch-ai folder and install the project requirements. Create a MySQL database named devmatch_ai and an application user with access to that database. Copy .env.example to .env and set DATABASE_URL to the MySQL connection URL and SECRET_KEY to a private random value. Keep credentials outside source code.')
code('python -m venv .venv\n.venv\\Scripts\\Activate.ps1\npip install -r requirements.txt\nflask --app app init-db\nflask --app app seed-demo\nflask --app app verify-dbms\nflask --app app run --debug')
para('Open the local Flask address in a browser. Run mysql_demo.sql in the same MySQL database to create the summary view, audit table and status trigger. The seed command supplies demonstration accounts and records; it should be used in a dedicated demo database. The verifier checks table rows, attributes and relationships.')
para('For hosted execution, set production configuration, approved hosts, a reachable MySQL database, private Supabase Storage credentials and shared Redis rate limiting. The app validates required production settings at startup. Gemini is optional; its API key and model setting remain server side.')
newpage();heading('4.2 Front End Code')
para('The front end is organized as Jinja templates extending base.html. Job cards are rendered from the current query results. Each application form includes a CSRF token, and developer only actions are conditionally displayed. The following excerpts come from jobs.html, with line wrapping adjusted for the report.')
code('''{% for job in jobs %}
<article class="job-row">
  <div class="small-muted">
    {{job.recruiter.company_name}} · {{job.location}}
  </div>
  <h2>{{job.title}}</h2>
  <p>{{job.description}}</p>
  {% for skill in job.required_skills.split(',') %}
    <span>{{skill.strip()}}</span>
  {% endfor %}
</article>
{% endfor %}''')
para('The apply form submits to the Flask route for the selected job. The server rechecks identity, eligibility and existing applications; conditional rendering alone is not used as the access control.')
code('''<form method="post"
      action="{{url_for('apply',id=job.id)}}">
  <input type="hidden" name="csrf" value="{{csrf}}">
  <button class="button small">Apply now ↗</button>
</form>''')
para('The developer dashboard summarizes skills, projects and applications, while the recruiter dashboard exposes job management and applicant review. Search forms use query parameters for role or skill and location. JavaScript enhances interaction, and responsive CSS adapts the page layout to available screen width.')
newpage();heading('4.3 Back End Code')
para('The backend defines the ten application tables as SQLAlchemy models. Password hashing is used for credentials. Session checks restrict authenticated workflows, and role checks protect developer and recruiter routes. Record ownership is checked before profile, job, skill or application changes.')
heading('Account and portfolio modules')
code('''class User(db.Model):
    __tablename__ = 'users'
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), nullable=False)
    email = db.Column(db.String(190), unique=True,
                      nullable=False, index=True)
    password_hash = db.Column(db.String(255),
                              nullable=False)
    role = db.Column(db.String(20), nullable=False)''')
para('Registration validates account information and creates the role specific profile. Login checks the password hash and stores the authenticated user in the session. The developer portfolio module manages profile fields, skills and projects; technologies entered in a form are stored in project_technologies child rows. Recruiter profile fields describe the company rather than duplicating the shared credential record.')
heading('Application model and uniqueness')
code('''class Application(db.Model):
    __tablename__ = 'applications'
    id = db.Column(db.Integer, primary_key=True)
    developer_id = db.Column(db.Integer,
        db.ForeignKey('developer_profiles.id'), nullable=False)
    job_id = db.Column(db.Integer,
        db.ForeignKey('jobs.id'), nullable=False)
    status = db.Column(db.String(30), default='Applied',
                       nullable=False)
    created_at = db.Column(db.DateTime,
        default=lambda: datetime.now(timezone.utc))
    __table_args__ = (UniqueConstraint('developer_id',
        'job_id', name='uq_application'),)''')
newpage();heading('4.3 Back End Code continued')
para('The job module allows recruiters to create, edit and remove their own postings. Required skill values are stored as job_skills child records. The application module checks whether the job exists, whether its deadline has passed and whether the developer already applied. A successful submission creates an application and commits it to MySQL.')
heading('Skill matching and initial status')
para('Let R be the set of case normalized job skill names and C the union of profile skill names and the latest extracted resume skills. The displayed score is round(100 × |R ∩ C| / |R|), or zero when R is empty. The shortlist decision uses the unrounded fraction and a strict threshold above 70 percent.')
code('''SHORTLIST_MATCH_THRESHOLD = 70

def should_shortlist(matched_count, required_count):
    return (required_count > 0 and
        100 * matched_count >
        SHORTLIST_MATCH_THRESHOLD * required_count)

# Within the application submission route
developer = current_user().developer
result = ats_match(developer, job)
status_value = ('Shortlisted' if
    result['recommend_shortlist'] else 'Rejected')
db.session.add(Application(job_id=id,
    developer_id=developer.id, status=status_value))
db.session.commit()''')
para('A candidate matching three of three skills scores 100 percent and meets the threshold; two of three scores 67 percent and does not. Exactly seven of ten skills produces 70 percent and is rejected initially. A job with no recorded skills also does not meet the threshold. The recruiter can change the status to Applied, Under Review, Shortlisted, Rejected or Selected after review.')
para('The function exposes matched, missing, profile matched and resume matched skill lists. These lists explain the score. Matching is based on recorded names, not semantic equivalence, and does not use proficiency levels, years, minimum experience or protected personal characteristics.')
newpage();heading('4.3 Back End Code continued')
heading('Resume analysis and storage modules')
para('The resume route accepts a PDF within the configured 5 MB request limit, sanitizes the filename and extracts document text using PyMuPDF. Structured analysis is saved as a ResumeAnalysis row linked to the developer. The matching function uses the latest analysis rather than combining every historical upload.')
code('''# Actual application configuration
app.config.update(
    SQLALCHEMY_DATABASE_URI=DATABASE_URL,
    SQLALCHEMY_TRACK_MODIFICATIONS=False,
    MAX_CONTENT_LENGTH=5 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=IS_PRODUCTION,
)''')
para('When a Gemini API key is configured, extracted resume text is submitted for structured analysis. Without a key or when the analysis request fails, the local extraction path identifies supported skills and returns a deterministic result. The career advisor uses the developer’s profile and current job listings as context. It requires the configured AI service.')
para('Local development can store PDF files in the upload directory. Production configuration uses a private Supabase bucket, and the authenticated resume route retrieves the file. MySQL stores the application records and resume path; Supabase’s PostgreSQL database is not used by this project.')
heading('Security and operational controls')
para('The source includes CSRF checks, role and ownership checks, password hashing, login and registration rate limits, secure production cookies and response security headers. Production startup requires secret configuration, allowed hosts, private storage and shared rate limit storage. Secrets are server configuration values and must not appear in templates or screenshots. These controls support the current demonstration; broader upload scanning and operational monitoring remain future work.')
newpage();heading('4.3 Back End Code continued')
heading('SQL queries view and trigger')
para('The supplied mysql_demo.sql demonstrates joins, grouping, filtering, a view and status auditing. The following view combines job ownership with application counts, including jobs with no applicants because the application join is a LEFT JOIN.')
code('''CREATE OR REPLACE VIEW job_application_summary AS
SELECT j.id AS job_id, j.title, r.company_name,
       COUNT(a.id) AS applicant_count,
       SUM(CASE WHEN a.status = 'Shortlisted'
                THEN 1 ELSE 0 END) AS shortlisted_count
FROM jobs j
JOIN recruiter_profiles r ON r.id = j.recruiter_id
LEFT JOIN applications a ON a.job_id = j.id
GROUP BY j.id, j.title, r.company_name;''')
para('The status trigger creates an audit row only when an UPDATE changes the application status. The SQL script separately backfills initial status records for existing applications. It does not install an INSERT trigger for every future application.')
code('''DELIMITER $$
CREATE TRIGGER log_application_status
AFTER UPDATE ON applications FOR EACH ROW
BEGIN
  IF OLD.status <> NEW.status THEN
    INSERT INTO application_status_history
      (application_id, old_status, new_status)
    VALUES (NEW.id, OLD.status, NEW.status);
  END IF;
END$$
DELIMITER ;''')
para('The script also uses ALTER TABLE to widen the salary column and demonstrates application INSERT, UPDATE and DELETE in a transaction ending in ROLLBACK. This keeps the transaction’s test application from persisting. Website CRUD actions commit their changes; a new request reads updated records. The /dbms-demo page queries MySQL again and disables browser caching.')

# 16-18 Results snapshots
newpage(5,'Interpretation of Result')
para('The following screenshots show the locally running application connected to MySQL. They illustrate public navigation, portfolio records, job matching and recruiter review. Figure captions and interpretations relate each interface to its stored data.')
heading('5.1 Result Snapshots')
picture('home',5.4);caption('Fig 5.1 Public home page')
para('The home page in Fig 5.1 provides job search and account entry points. The displayed vacancy count comes from the database. Visitors can browse opportunities before logging in, while application submission requires a developer account.',11)
picture('developer',5.4);caption('Fig 5.2 Developer dashboard')
para('Fig 5.2 shows a seeded developer’s skills, project and application. The Shortlisted status is read from MySQL. Portfolio child tables are linked to the developer profile.',11)
newpage();heading('5.1 Result Snapshots continued')
picture('match',5.4);caption('Fig 5.3 Job listing and skill match')
para('The job listing in Fig 5.3 connects each vacancy to company details and presents the developer’s skill coverage. Matching uses profile and latest resume skills. Search and filtering operate on database backed vacancy records, and existing applications are marked in the interface.',11)
picture('applicants',5.4);caption('Fig 5.4 Recruiter applicant review')
para('Fig 5.4 displays a 100 percent match for Python, Flask and MySQL, with all three supplied by the profile. The interface explicitly indicates that the resume has not been analyzed. The recruiter sees matched and missing evidence and can change the decision through the status control.',11)
newpage();heading('5.1 Result Snapshots continued')
picture('dbms',5.4);caption('Fig 5.5 MySQL jobs demonstration')
para('Fig 5.5 displays current jobs queried from MySQL. This page supports comparing SQL records with the website. The schema verifier was executed against the local database and passed for all ten application tables. Table 5.1 records the observed counts; the audit extension was present too.',11)
table('Table 5.1 Database verification results',['Table','Rows','Attributes','Related'],[('users',12,6,'Yes'),('developer_profiles',6,10,'Yes'),('recruiter_profiles',6,6,'Yes'),('skills',15,5,'Yes'),('projects',5,6,'Yes'),('project_technologies',10,6,'Yes'),('jobs',15,10,'Yes'),('job_skills',27,6,'Yes'),('applications',10,5,'Yes'),('resume_analysis',6,7,'Yes'),('application_status_history',7,5,'Yes')],[7.8,2.1,2.8,2.5])
para('The observed result confirms the row, attribute and relationship minimums. Developer and recruiter login, dashboards and job pages loaded during screenshot capture. These observations do not establish production capacity or semantic accuracy of AI analysis.',11)

# 19 conclusion
newpage(6,'Conclusion and Future Scope')
para('DevMatch AI implements a relational developer portfolio and hiring workflow with separate developer and recruiter accounts. The database connects profiles, skills, projects, vacancies, applications and resume analyses. The interface gives users a consistent way to maintain their information and gives recruiters a visible record of candidate skill matches and application decisions.')
para('The local schema verification passed the required minimums for the ten application tables. The SQL audit extension was also present with related history records. The project demonstrates primary and foreign keys, uniqueness constraints, normalized repeating values, CRUD, joins, filtering, grouping, a summary view and a status change trigger. The website and SQL demonstration query the same MySQL database.')
heading('Learning outcomes')
para('The project develops practical understanding of ER modeling, mapping cardinalities to foreign keys and bridge tables, removing multivalued attributes and connecting database transactions to user workflows. It also illustrates why ownership checks and constraints must support interface validation. Explainable matching shows how a computed result can be connected to stored evidence.')
heading('Future scope')
para('Future work can introduce schema migration tooling, automated regression tests, concurrency handling for duplicate submissions and a more complete event history. Resume analysis can add OCR for scanned documents, a controlled skill vocabulary and explicit evaluation against labeled examples. Matching can account for verified proficiency and job priorities, while keeping recruiter review available.')
para('Further extensions include interview scheduling, notifications, analytics and accessibility improvements. Public deployment would benefit from file scanning, backups, monitoring and resume retention controls. MySQL remains the application database.')

# 20 references
newpage();p=para('REFERENCES',16,True,WD_ALIGN_PARAGRAPH.CENTER);p.style=doc.styles['Heading 1']
para('The following technical references describe the database, framework and optional integrations used in this project. The project’s own source and documentation identify its specific behavior and schema.')
refs=[('Oracle','MySQL 8.4 Reference Manual','https://dev.mysql.com/doc/refman/8.4/en/'),('Pallets','Flask Documentation','https://flask.palletsprojects.com/en/stable/'),('SQLAlchemy authors','SQLAlchemy 2.0 ORM Documentation','https://docs.sqlalchemy.org/en/20/orm/'),('Artifex','PyMuPDF Documentation','https://pymupdf.readthedocs.io/en/latest/'),('Google','Gemini API Documentation','https://ai.google.dev/gemini-api/docs'),('Supabase','Storage Documentation','https://supabase.com/docs/guides/storage')]
for i,(a,t,u) in enumerate(refs,1):para(f'[{i}] {a}, {t}. [Online]. Available: {u}. Accessed: 3 October 2026.',12)
para('[7] DevMatch AI project source, app.py, templates, requirements.txt and mysql_demo.sql, local project repository.',12)
para('[8] DevMatch AI DBMS Project Documentation and README, local project repository.',12)
para('[9] CMR Institute of Technology, DBMS Mini Project Report Content Format and Cover Certificate Declaration Abstract Acknowledgement and TOC Format, supplied report templates.',12)

doc.core_properties.title='DevMatch AI DBMS Mini Project Report';doc.core_properties.subject='Developer portfolio and hiring platform';doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
settings=doc.settings.element
up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');settings.append(up)
final=OUT/'DevMatch_AI_DBMS_Mini_Project_Report.docx';doc.save(final)
assert len(doc.sections)==2
assert all(sha256(Path(f).read_bytes()).hexdigest()==h for f,h in sources.items())
assert len(replacements[49].split())<=100
print('Created',final,'planned body pages',page_index)
