from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(r'D:\DBMS_DevApp')
target=root/'output/DevMatch_AI_Final_Submission_Report.docx'
d=Document(target)
ref=Document(r'C:\Users\gurra\Downloads\DBMS Mini Project _Report content format.docx')
front=Document(r'C:\Users\gurra\Downloads\DBMS Mini Project_ cover+certificate+Declaration+Abstract+acknowledgement+TOC.docx')
title='DevMatch AI \u2013 Developer Portfolio & Hiring Platform'

# Keep the original inline logo intact, with an explicit right tab instead
# of the template's space padding, which breaks with a longer project title.
header=d.sections[1].header.paragraphs[0]
drawing=deepcopy(header._p.xpath('.//w:r[w:drawing]')[0])
ppr=deepcopy(ref.sections[0].header.paragraphs[0]._p.pPr)
header.clear()
if header._p.pPr is not None:header._p.remove(header._p.pPr)
header._p.insert(0,ppr)
for el in header._p.pPr.xpath('./w:tabs'):header._p.pPr.remove(el)
tabs=OxmlElement('w:tabs')
for pos in (4513,4680,9026,9360):
 el=OxmlElement('w:tab');el.set(qn('w:val'),'clear');el.set(qn('w:pos'),str(pos));tabs.append(el)
header._p.pPr.append(tabs)
width=d.sections[1].page_width-d.sections[1].left_margin-d.sections[1].right_margin
header.paragraph_format.tab_stops.add_tab_stop(width,WD_TAB_ALIGNMENT.RIGHT)
header.paragraph_format.line_spacing=1.0
r=header.add_run(title+'\t');r.font.name='Times New Roman';r.font.size=Pt(9);r.italic=True
header._p.append(drawing)

# Match the footer's three positions: department, centered year, right page.
footer=d.sections[1].footer.paragraphs[0]
ppr=deepcopy(ref.sections[0].footer.paragraphs[0]._p.pPr)
footer.clear()
if footer._p.pPr is not None:footer._p.remove(footer._p.pPr)
footer._p.insert(0,ppr)
for el in footer._p.pPr.xpath('./w:tabs'):footer._p.pPr.remove(el)
footer.paragraph_format.tab_stops.add_tab_stop(int(width/2),WD_TAB_ALIGNMENT.CENTER)
footer.paragraph_format.tab_stops.add_tab_stop(width,WD_TAB_ALIGNMENT.RIGHT)
footer.paragraph_format.line_spacing=1.0
footer.paragraph_format.space_after=Pt(0)
r=footer.add_run('Dept. of ECE, CMRIT\t2026-2027\tPage ')
r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=True
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGE ')
rxml=OxmlElement('w:r');rp=deepcopy(r._r.rPr);rxml.append(rp)
t=OxmlElement('w:t');t.text='1';rxml.append(t);field.append(rxml);footer._p.append(field)

paras=d.paragraphs
for p in paras:
 if p.text.startswith('Chapter '):
  p.runs[0].text=p.runs[0].text.replace('Chapter ','CHAPTER ')
 if p.text in ('Acknowledgement','ACKNOWLEDGEMENT'):
  p.paragraph_format.space_before=Pt(65)
  p.paragraph_format.space_after=Pt(35)
 if p.text.startswith('We express our gratitude') or p.text.startswith('We thank'):
  p.paragraph_format.space_after=Pt(12)
 if p.text=='LIST OF TABLES':p.paragraph_format.space_before=Pt(130)

# Restore the institutional colour accents in the retained front template.
for i in (3,4,12,13,15,23):
 colour='0000FF' if i in (3,4) else '0070C0' if i in (12,13) else 'C00000'
 for r in paras[i].runs:r.font.color.rgb=RGBColor.from_string(colour)

# Source signature lines are centered and the role labels are emphasized.
for c in (d.tables[0].cell(0,0),d.tables[0].cell(0,2)):
 for p in c.paragraphs:
  p.text=p.text.replace('Name ____________________','Name ______________').replace('Designation ____________________','Designation __________')
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=True

# Restore the black unshaded contents grid; figures/tables are borderless
# aligned lists, as shown in the front-matter template.
for index in (2,3,4):
 table=d.tables[index]
 for row in table.rows:
  for j,c in enumerate(row.cells):
   pr=c._tc.get_or_add_tcPr()
   for el in pr.xpath('./w:shd | ./w:tcBorders'):pr.remove(el)
   borders=OxmlElement('w:tcBorders')
   for side in ('top','left','bottom','right'):
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single' if index==2 else 'nil')
    el.set(qn('w:sz'),'4');el.set(qn('w:color'),'000000');borders.append(el)
   pr.append(borders)
   for p in c.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==1 else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11)
 # Remove inherited table-level borders/fills too.
 pr=table._tbl.tblPr
 for el in pr.xpath('./w:tblBorders | ./w:shd'):pr.remove(el)
 if index in (3,4):
  for row in table.rows[1:]:
   for p in row.cells[0].paragraphs:
    for r in p.runs:r.bold=True

for el in d.settings.element.xpath('./w:updateFields'):el.set(qn('w:val'),'false')
for borders in d.sections[0]._sectPr.xpath('./w:pgBorders'):
 borders.set(qn('w:display'),'allPages');borders.set(qn('w:zOrder'),'back')
 for edge in borders:edge.set(qn('w:color'),'000000')
d.save(target)
print('Formatting corrected:',target)
