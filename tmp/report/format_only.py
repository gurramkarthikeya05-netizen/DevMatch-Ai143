from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

root=Path(r'D:\DBMS_DevApp');f=root/'output/DevMatch_AI_Final_Submission_Report.docx'
d=Document(f);body=d._element.body
def delete(p):
 if p._p.getparent() is not None:p._p.getparent().remove(p._p)
def rewrite(p,text):
 p.clear();r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor(0,0,0)
 return p
remove_prefixes=[
 'The local implementation uses MySQL and can store resume files locally.',
 'For hosted execution, set production configuration',
 'First normal form is supported by individual skill and technology rows.',
 'The source includes CSRF checks, role and ownership checks,',
 '[7]','[8]','[9]']
remove_labels={'Normalization and integrity','Security and operational controls','Learning outcomes','Future scope'}
for p in list(d.paragraphs):
 if p.text in remove_labels or any(p.text.startswith(s) for s in remove_prefixes):delete(p)
 elif p.text in ('Account and portfolio modules','Application model and uniqueness','Skill matching and initial status','Resume analysis and storage modules','SQL queries view and trigger'):
  p.style=d.styles['Normal']
 elif p.text=='4.2 Front End Code':rewrite(p,'4.2 Front-end Code');p.runs[0].bold=True;p.runs[0].font.size=Pt(14)
 elif p.text=='4.3 Back End Code':rewrite(p,'4.3 Back-end Code');p.runs[0].bold=True;p.runs[0].font.size=Pt(14)
 elif p.text=='4.3 Back End Code continued':rewrite(p,'4.3 Back-end Code continued');p.runs[0].bold=True;p.runs[0].font.size=Pt(14)
 elif p.text.startswith('The job_id and name pair is unique in job_skills.'):
  rewrite(p,'The job_id and name pair is unique in job_skills. Each job is linked to its recruiter through recruiter_id, and each requirement is linked to its job through job_id. These relationships support job management and skill matching without storing comma separated skill values in the jobs table.')

# Required implementation modules remain within subsection 4.3.
for p in d.paragraphs:
 if p.text.startswith('class User(db.Model):'):
  p.clear();r=p.add_run("""# Login module excerpt
user = User.query.filter_by(email=email).first()
if user and check_password_hash(user.password_hash, password):
    session.clear()
    session['uid'] = user.id
    return redirect(url_for('dashboard'))

# Profile module excerpt
profile = (user.developer if user.role == 'developer'
           else user.recruiter)
for field in fields:
    value = request.form.get(field, '').strip()
    setattr(profile, field, value)
db.session.commit()""");r.font.name='Consolas';r.font.size=Pt(9)
 elif p.text.startswith('SHORTLIST_MATCH_THRESHOLD = 70'):
  p.clear();r=p.add_run("""# Skill and project module excerpts
db.session.add(Skill(developer_id=current_user().developer.id,
    name=name, level=request.form.get('level','Intermediate'),
    years=max(0,min(50,float(request.form.get('years') or 0)))))
db.session.commit()
project.technology_items = [ProjectTechnology(name=technology)
    for technology in parse_list_values(
        request.form.get('technologies',''))]
db.session.add(project)
db.session.commit()

# Matching and application module excerpt
def should_shortlist(matched_count, required_count):
    return (required_count > 0 and
        100 * matched_count > 70 * required_count)
result = ats_match(developer, job)
status_value = ('Shortlisted' if
    result['recommend_shortlist'] else 'Rejected')
db.session.add(Application(job_id=id,
    developer_id=developer.id, status=status_value))
db.session.commit()""");r.font.name='Consolas';r.font.size=Pt(9)
 elif p.text.startswith('# Actual application configuration'):
  p.clear();r=p.add_run("""# Resume analysis module excerpt
document = pymupdf.open(stream=data, filetype='pdf')
text = '\\n'.join(page.get_text() for page in document)[:25000]
analysis = analyze_resume(text)
db.session.add(ResumeAnalysis(developer_id=developer.id,
    summary=analysis['summary'],
    strengths=', '.join(analysis['strengths']),
    missing_skills=', '.join(analysis['missing_skills']),
    extracted_skills=', '.join(analysis['skills'])))
db.session.commit()

# Job management module excerpt
job.skills = [JobSkill(name=skill) for skill in required]
db.session.add(job)
db.session.commit()
job = Job.query.filter_by(id=id,
    recruiter_id=current_user().recruiter.id).first_or_404()
db.session.delete(job)
db.session.commit()""");r.font.name='Consolas';r.font.size=Pt(9)

# Replace the result material with precisely the three snapshot slots shown
# in the supplied content format, each followed by its interpretation.
paras=list(d.paragraphs)
start=next(p for p in paras if p.text=='Chapter 5')._p
end=next(p for p in paras if p.text=='Chapter 6')._p
elements=list(body);a=elements.index(start);b=elements.index(end)
saved=elements[a:a+4]
intro='This chapter presents the main results of DevMatch AI. The snapshots show developer portfolio management, job search and skill matching, and recruiter review of applications. Each figure is followed by an interpretation of the displayed result.'
chapter_intro=next(p for p in paras if p.text.startswith('The following screenshots show'))
rewrite(chapter_intro,intro)
heading=next(p for p in paras if p.text=='5.1 Result Snapshots')
saved=[start,start.getnext(),chapter_intro._p,heading._p]
for el in elements[a:b]:body.remove(el)
for el in saved:body.insert(list(body).index(end),el)
def append_before_end(text,size=12,center=False,break_before=False):
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
 p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(8);p.paragraph_format.page_break_before=break_before
 r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0)
 body.remove(p._p);body.insert(list(body).index(end),p._p);return p
snapshots=[
 ('developer', 'Fig 5.1 Developer dashboard', 'The developer dashboard displays the stored profile, skills, projects and applications. The sample developer has three skills, one project and one application. The application status is displayed alongside the job title. Skills and projects are stored in child tables linked to the developer profile, while resume analysis is linked separately. This result shows that portfolio information and application tracking are brought together through relational database queries.'),
 ('match', 'Fig 5.2 Job listing and skill match', 'The job page displays vacancies with company details, location, required skills and the developer’s skill match. Search fields filter the current database records. The match is calculated from the intersection of recorded job skills with profile skills and the latest extracted resume skills. The application action depends on login, the closing date and whether an application already exists. This result connects the jobs, job_skills, recruiter_profiles and applications tables in one workflow.'),
 ('applicants', 'Fig 5.3 Recruiter applicant review', 'The recruiter dashboard displays applicants for the recruiter’s own jobs. The illustrated candidate matches Python, Flask and MySQL, giving a 100 percent skill match. The screen identifies the profile as the evidence source and states that the resume has not been analyzed. Matched and missing skills help the recruiter interpret the score. The status control allows the recruiter to review and revise the application decision, which is saved to the applications table.')]
for i,(image,cap,interpretation) in enumerate(snapshots):
 p=append_before_end('',break_before=i>0,center=True);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=True
 p.add_run().add_picture(str(root/'tmp/report'/(image+'.png')),width=Inches(5.7))
 p=append_before_end(cap,10.5,True);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=True
 p=append_before_end('Interpretation: '+interpretation)
 p.runs[0].text=interpretation
 r=p.add_run('')
 # Preserve a separate bold interpretation label as in the reference.
 p.clear();r=p.add_run('Interpretation: ');r.bold=True;r.font.name='Times New Roman';r.font.size=Pt(12)
 r=p.add_run(interpretation);r.font.name='Times New Roman';r.font.size=Pt(12)

# Match the TOC entries and figure/table inventories to the format sections.
toc=d.tables[2]
for row in toc.rows:
 tx=row.cells[0].text
 if tx=='3.3 Database Tables and Normalization':row.cells[0].text='3.3 Database Tables'
 if tx=='4.2 Front End Code':row.cells[0].text='4.2 Front-end Code'
 if tx=='4.3 Back End Code and Modules':row.cells[0].text='4.3 Back-end Code'
fig=d.tables[3]
while len(fig.rows)>1:fig._tbl.remove(fig.rows[-1]._tr)
for tx,page in [('Fig 3.1 Relational schema','4'),('Fig 3.2 Entity relationships and cardinalities','5'),('Fig 5.1 Developer dashboard','16'),('Fig 5.2 Job listing and skill match','17'),('Fig 5.3 Recruiter applicant review','18')]:
 cells=fig.add_row().cells;cells[0].text=tx;cells[1].text=page
tabs=d.tables[4]
for row in list(tabs.rows):
 if 'Table 5.1' in row.cells[0].text:tabs._tbl.remove(row._tr)
# Restore table formatting for updated front rows without changing page geometry.
for t in (toc,fig,tabs):
 for i,row in enumerate(t.rows):
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=i==0
for el in d.settings.element.xpath('./w:updateFields'):el.set(qn('w:val'),'false')
d.save(f)
print('Saved revised report:',f)
